from dataclasses import dataclass
from pathlib import Path

from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

CHUNK_WORDS = 420
CHUNK_OVERLAP = 70
RESULT_COUNT = 4


@dataclass(frozen=True)
class Passage:
    source: str
    group: str
    text: str


def _extract_text(path: Path) -> str:
    document = Document(path)
    sections = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            sections.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(section for section in sections if section)


def _split_into_chunks(text: str) -> list[str]:
    words = text.split()
    if not words:
        return []

    step = CHUNK_WORDS - CHUNK_OVERLAP
    return [
        " ".join(words[start : start + CHUNK_WORDS])
        for start in range(0, len(words), step)
    ]


class KnowledgeBase:
    def __init__(self, root: Path) -> None:
        if not root.is_dir():
            raise FileNotFoundError(f"Knowledge folder was not found: {root}")

        passages = []
        for path in sorted(root.rglob("*.docx")):
            text = _extract_text(path)
            relative_path = path.relative_to(root)
            group = (
                relative_path.parts[0]
                if len(relative_path.parts) > 1
                else "general"
            )
            passages.extend(
                Passage(source=path.name, group=group, text=chunk)
                for chunk in _split_into_chunks(text)
            )

        if not passages:
            raise FileNotFoundError(
                "No readable DOCX documents found under "
                f"{root}. Add knowledge documents before starting the API."
            )

        self._passages = passages
        self._vectorizer = TfidfVectorizer(
            ngram_range=(1, 2),
            strip_accents="unicode",
            sublinear_tf=True,
            max_features=100_000,
        )
        self._matrix = self._vectorizer.fit_transform(
            [passage.text for passage in passages]
        )

    def search(self, query: str) -> list[Passage]:
        query_vector = self._vectorizer.transform([query])
        scores = cosine_similarity(query_vector, self._matrix).ravel()
        ranked_indices = scores.argsort()[::-1]
        return [
            self._passages[index]
            for index in ranked_indices[:RESULT_COUNT]
            if scores[index] > 0
        ]
