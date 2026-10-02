import tempfile
import unittest
from pathlib import Path

from docx import Document

from backend.knowledge_base import KnowledgeBase


class KnowledgeBaseTests(unittest.TestCase):
    def test_search_returns_matching_document_and_folder(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            research_folder = root / "research"
            research_folder.mkdir()

            document = Document()
            document.add_paragraph(
                "The fictional Zorblax permit requires annual renewal."
            )
            document.save(research_folder / "permits.docx")

            knowledge_base = KnowledgeBase(root)
            results = knowledge_base.search(
                "When must the Zorblax permit renew?"
            )

        self.assertTrue(results)
        self.assertEqual(results[0].source, "permits.docx")
        self.assertEqual(results[0].group, "research")
        self.assertFalse(knowledge_base.search("volcanic rock samples"))

    def test_search_reads_table_content(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            project_folder = root / "nokshi_threads"
            project_folder.mkdir()

            document = Document()
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Nimbusk tax record"
            table.cell(0, 1).text = "Renewal is due each spring."
            document.save(project_folder / "records.docx")

            knowledge_base = KnowledgeBase(root)
            results = knowledge_base.search("Nimbusk tax record renewal")

        self.assertTrue(results)
        self.assertEqual(results[0].source, "records.docx")
        self.assertEqual(results[0].group, "nokshi_threads")

    def test_empty_folder_has_actionable_error(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            with self.assertRaisesRegex(FileNotFoundError, "No readable DOCX"):
                KnowledgeBase(Path(temporary_directory))


if __name__ == "__main__":
    unittest.main()
